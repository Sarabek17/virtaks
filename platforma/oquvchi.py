# -*- coding: utf-8 -*-
"""Manba o'quvchilari — istalgan formatni bo'laklarga aylantiradi.

Har o'quvchi bir xil natija qaytaradi:
    {"bolaklar": [{joy, matn, audio_bosh?, audio_oxir?, sahifa?}, ...],
     "davomiylik": int|None,        # audio soniya
     "rasmlar": [(sahifa_idx, bayt, mime), ...]}   # slayd/sahifa suratlari

Bo'lakdagi `audio_bosh/audio_oxir` va `sahifa` keyin AYNAN OLINGAN FRAGMENTNI
qaytarish uchun ishlatiladi (dars yozuvining o'sha daqiqasi / o'sha slayd rasmi).

Og'ir ish (STT, OCR, LibreOffice) faqat WORKER'da bajariladi.
"""
import contextvars
import re
import shutil
import subprocess
import tempfile
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime
from pathlib import Path

from google.genai import types

from . import llm
from .sozlama import log

AUDIO_TURLAR = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac",
                ".mp4", ".mkv", ".avi", ".mov", ".webm"}
SLAYD_TURLAR = {".pdf", ".pptx", ".ppt"}
JADVAL_TURLAR = {".xlsx", ".xlsm"}
HUJJAT_TURLAR = {".docx", ".txt", ".md"}

MAKS_BOLAK = 1200        # belgi
KATAK_UZUNLIK = 200      # jadval katagidagi matn chegarasi
FORMULA_MAKS_MB = 20     # bundan katta xlsx'da formula o'qish o'tkazib yuboriladi
STT_BOLAK_DAQIQA = 10    # audio nechchi daqiqalik qismlarga bo'linadi
PARALLEL = 4
RASM_MASSHTAB = 2.0      # 144 dpi
JPEG_SIFAT = 80

STT_PROMPT = """Sen professional o'zbek tili transkripsiya mutaxassisisan. Ushbu audiodagi BARCHA gaplarni so'zma-so'z yozib chiq.

QOIDALAR:
1. O'zbek lotin alifbosida yoz.
2. So'zlovchi qanday gapirsa, SHUNDAY yoz — adabiy tilga "tuzatma". Sheva so'zlarini (qanaqa, shetta, bo'ganda, qiganda...) aslidek qoldir.
3. Ruscha so'zlarni aslidek qoldir (lyuboy, srazu, chto-to, uje...).
4. Arabcha oyat/duo bo'lsa, lotin transliteratsiyada yoz.
5. Tushunarsiz joyni [tushunarsiz] deb belgila.
6. Tinish belgilarini to'g'ri qo'y, har gapni bosh harf bilan boshla.
7. VAQT BELGISI: har abzats boshiga audiodagi vaqtni [MM:SS] ko'rinishida yoz (masalan [03:45]). Har 2-4 gapda yangi abzats boshla. Vaqtni shu audio bo'lagining boshidan hisobla ([00:00] dan boshlanadi).
8. FAQAT transkript matnini chiqar — hech qanday izoh, sarlavha, xulosa yozma."""

OCR_PROMPT = """Bu o'quv materiali sahifasining (slayd yoki hujjat) rasmi. Undagi BARCHA mazmunni to'liq Markdown formatida chiqar.

QOIDALAR:
1. Sahifa sarlavhasini `#` bilan boshla.
2. Barcha matnni aslidek yoz (o'zbek/rus/ingliz — qanday bo'lsa shunday). Hech narsani tashlab ketma.
3. Jadval bo'lsa — Markdown jadval qil.
4. Diagramma, sxema, rasm bo'lsa — mazmunini 1-3 jumlada tavsifla: `[RASM: ...]`.
5. Sahifa raqami, kolontitul, suv belgisi kabi bezaklarni YOZMA.
6. FAQAT sahifa mazmunini chiqar — o'z izohingni qo'shma."""


# ---------------------------------------------------------------- yordamchilar

def tur_aniqla(nom: str) -> str:
    """Fayl nomidan manba turini aniqlaydi."""
    s = Path(nom).suffix.lower()
    if s in AUDIO_TURLAR:
        return "audio"
    if s in SLAYD_TURLAR:
        return "slayd"
    if s in JADVAL_TURLAR:
        return "jadval"
    if s in HUJJAT_TURLAR:
        return "hujjat"
    return ""


def ffmpeg_yol() -> str:
    """Tizimdagi ffmpeg (Docker image'da bor), bo'lmasa imageio-ffmpeg (lokal)."""
    yol = shutil.which("ffmpeg")
    if yol:
        return yol
    import imageio_ffmpeg
    return imageio_ffmpeg.get_ffmpeg_exe()


def audio_davomiyligi(manba: str) -> float:
    """Sekundlarda (fayl yo'li yoki URL)."""
    r = subprocess.run([ffmpeg_yol(), "-i", str(manba)],
                       capture_output=True, text=True, errors="replace")
    m = re.search(r"Duration:\s*(\d+):(\d+):(\d+(?:\.\d+)?)", r.stderr)
    if not m:
        raise RuntimeError(f"Davomiylik aniqlanmadi: {str(manba)[:120]}")
    return int(m.group(1)) * 3600 + int(m.group(2)) * 60 + float(m.group(3))


def _vaqt(s: int) -> str:
    return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}"


# ---------------------------------------------------------------- bo'laklash

_VAQT_BELGI = re.compile(r"^\[(\d{1,2}):(\d{2}):(\d{2})\]")


def audio_bolaklar(matn: str, davomiylik: int) -> list[dict]:
    """Vaqt belgili transkriptni ~MAKS_BOLAK belgili bo'laklarga yig'adi.

    Har bo'lakka aniq vaqt oralig'i qo'yiladi: bosh — birinchi abzats vaqti,
    oxir — bo'lakdan KEYINGI abzats vaqti (gap o'rtasidan kesilmasin).
    """
    abzatslar = []
    for q in matn.split("\n"):
        q = q.strip()
        if not q:
            continue
        m = _VAQT_BELGI.match(q)
        t = (int(m.group(1)) * 3600 + int(m.group(2)) * 60 + int(m.group(3))) if m else None
        abzatslar.append((t, q))
    if not abzatslar:
        return []

    # guruhlash: MAKS_BOLAK belgiga yetganda yangi bo'lak (1 abzats ustma-ust)
    guruhlar, joriy = [], []
    for i, (_, q) in enumerate(abzatslar):
        joriy.append(i)
        if sum(len(abzatslar[j][1]) for j in joriy) >= MAKS_BOLAK:
            guruhlar.append(joriy)
            joriy = [i]
    if len(joriy) > 1 or (joriy and not guruhlar):
        guruhlar.append(joriy)

    bolaklar = []
    for g in guruhlar:
        vaqtlar = [abzatslar[j][0] for j in g if abzatslar[j][0] is not None]
        bosh = vaqtlar[0] if vaqtlar else 0
        keyingi = g[-1] + 1
        oxir = None
        while keyingi < len(abzatslar):
            if abzatslar[keyingi][0] is not None:
                oxir = abzatslar[keyingi][0]
                break
            keyingi += 1
        if oxir is None or oxir <= bosh:
            oxir = min(davomiylik or bosh + 90, bosh + 90)
        bolaklar.append({
            "joy": f"[{_vaqt(bosh)}–{_vaqt(oxir)}]",
            "matn": "\n".join(abzatslar[j][1] for j in g),
            "audio_bosh": bosh, "audio_oxir": oxir})
    return bolaklar


def matn_bolaklar(matn: str, prefiks: str = "qism") -> list[dict]:
    abzatslar = [a.strip() for a in matn.split("\n\n") if a.strip()]
    bolaklar, joriy = [], []
    for a in abzatslar:
        joriy.append(a)
        if sum(len(x) for x in joriy) >= MAKS_BOLAK:
            bolaklar.append({"joy": f"{len(bolaklar) + 1}-{prefiks}",
                             "matn": "\n\n".join(joriy)})
            joriy = []
    if joriy:
        bolaklar.append({"joy": f"{len(bolaklar) + 1}-{prefiks}",
                         "matn": "\n\n".join(joriy)})
    return bolaklar


# ---------------------------------------------------------------- audio (STT)

def _stt_bolaklash(fayl: Path, papka: Path) -> list[Path]:
    """Audio/videoni STT_BOLAK_DAQIQA daqiqalik mp3 qismlarga bo'ladi."""
    papka.mkdir(parents=True, exist_ok=True)
    if fayl.suffix.lower() == ".mp3":
        kodlash = ["-c", "copy"]
    else:
        kodlash = ["-vn", "-ac", "1", "-ar", "16000", "-c:a", "libmp3lame", "-b:a", "64k"]
    r = subprocess.run(
        [ffmpeg_yol(), "-y", "-i", str(fayl), *kodlash,
         "-f", "segment", "-segment_time", str(STT_BOLAK_DAQIQA * 60),
         str(papka / "b_%03d.mp3")],
        capture_output=True, text=True, errors="replace")
    if r.returncode != 0:
        raise RuntimeError(f"ffmpeg xatosi: {r.stderr[-400:]}")
    return sorted(papka.glob("b_*.mp3"))


def _ofset_qoshish(matn: str, ofset_s: int) -> str:
    """Bo'lak ichidagi [MM:SS] belgilarni umumiy [H:MM:SS] vaqtga aylantiradi."""
    def alm(m):
        s = int(m.group(1)) * 60 + int(m.group(2)) + ofset_s
        return f"[{_vaqt(s)}]"
    return re.sub(r"\[(\d{1,2}):(\d{2})\]", alm, matn)


def _stt_qism(qism: Path, idx: int, jami: int, kesh, yoz) -> str:
    keshdan = kesh.ol(f"stt_{idx:03d}.txt")
    if keshdan is not None:
        yoz(f"  [{idx + 1}/{jami}] transkript keshdan ✓")
        return keshdan

    cl = llm.klient()
    yoz(f"  [{idx + 1}/{jami}] yuklanmoqda ({qism.stat().st_size / 1e6:.1f} MB)...")
    up = llm.qayta_urinib(lambda: cl.files.upload(file=str(qism)), f"yuklash {idx + 1}")
    while up.state and up.state.name == "PROCESSING":
        time.sleep(2)
        up = cl.files.get(name=up.name)
    try:
        matn = llm.generatsiya_qismlar(llm.STT_MODELLAR, [up, STT_PROMPT],
                                       harorat=0.2, maks_token=65535,
                                       bosqich="stt")
    finally:
        try:
            cl.files.delete(name=up.name)
        except Exception:
            pass
    matn = _ofset_qoshish(matn, idx * STT_BOLAK_DAQIQA * 60)
    kesh.yoz(f"stt_{idx:03d}.txt", matn)
    yoz(f"  [{idx + 1}/{jami}] transkript tayyor ✓ ({len(matn)} belgi)")
    return matn


def audio_oqi(fayl: Path, kesh, yoz) -> dict:
    dav = int(audio_davomiyligi(fayl))
    yoz(f"Audio: {fayl.name} | davomiylik {_vaqt(dav)}")
    with tempfile.TemporaryDirectory(prefix="twin_stt_") as tdir:
        qismlar = _stt_bolaklash(fayl, Path(tdir))
        yoz(f"{len(qismlar)} ta audio qism, parallel {PARALLEL}")
        natijalar: dict[int, str] = {}
        with ThreadPoolExecutor(max_workers=PARALLEL) as ex:
            # Xarajat konteksti oqimga NUSXA bilan uzatiladi (majlis.py:80
            # naqshi). Busiz STT sarfi `xarajatlar` jadvaliga UMUMAN
            # tushmasdi: ContextVar oqimni kesib o'tmaydi, `pul.xarajat_yoz`
            # esa kontekstsiz jimgina hech qayerga yozmasdi. B2B da bu
            # to'g'ridan-to'g'ri pul yo'qotish degani.
            fut = {ex.submit(contextvars.copy_context().run,
                             _stt_qism, q, i, len(qismlar), kesh, yoz): i
                   for i, q in enumerate(qismlar)}
            for f in as_completed(fut):
                natijalar[fut[f]] = f.result()
    matn = "\n\n".join(natijalar[i] for i in sorted(natijalar))
    return {"bolaklar": audio_bolaklar(matn, dav), "davomiylik": dav, "rasmlar": []}


# ---------------------------------------------------------------- PDF / PPTX

def pptx_pdfga(fayl: Path, papka: Path) -> Path:
    """LibreOffice headless orqali PPTX/PPT -> PDF (faqat worker image'da)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice topilmadi — PPTX faqat worker'da o'qiladi")
    r = subprocess.run(
        [soffice, "--headless", "--norestore", "--convert-to", "pdf",
         "--outdir", str(papka), str(fayl)],
        capture_output=True, text=True, errors="replace", timeout=600)
    pdf = papka / (fayl.stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError(f"PPTX->PDF bo'lmadi: {(r.stderr or r.stdout)[-300:]}")
    return pdf


OCR_MAKS_BELGI = 12000     # bitta slaydda bundan ko'p matn bo'lmaydi


def _sahifa_ocr(rasm: bytes, idx: int, jami: int, kesh, yoz) -> str:
    keshdan = kesh.ol(f"ocr_{idx:03d}.md")
    if keshdan is not None:
        yoz(f"  [{idx + 1}/{jami}] sahifa keshdan ✓")
        return keshdan
    qism = types.Part.from_bytes(data=rasm, mime_type="image/jpeg")
    matn = llm.generatsiya_qismlar(llm.OCR_MODELLAR, [qism, OCR_PROMPT],
                                   harorat=0.1, maks_token=8192, bosqich="ocr")
    # model ba'zan takrorlanish tsikliga tushadi (o'n minglab belgi) — bilim bazasini
    # ifloslantirmasligi uchun boshqa model bilan bir marta qayta, keyin kesamiz
    if len(matn) > OCR_MAKS_BELGI:
        yoz(f"  [{idx + 1}/{jami}] javob juda uzun ({len(matn)} belgi) — qayta o'qiladi")
        try:
            zanjir = llm.OCR_MODELLAR[1:] + llm.OCR_MODELLAR[:1]
            yangi = llm.generatsiya_qismlar(zanjir, [qism, OCR_PROMPT],
                                            harorat=0.0, maks_token=8192, bosqich="ocr")
            if len(yangi) < len(matn):
                matn = yangi
        except Exception as e:
            yoz(f"  [{idx + 1}/{jami}] qayta o'qish xatosi: {str(e)[:80]}")
        if len(matn) > OCR_MAKS_BELGI:
            matn = matn[:OCR_MAKS_BELGI] + "\n\n[...matn kesildi: model takrorlanib ketdi]"
    kesh.yoz(f"ocr_{idx:03d}.md", matn)
    yoz(f"  [{idx + 1}/{jami}] sahifa o'qildi ✓ ({len(matn)} belgi)")
    return matn


def pdf_oqi(fayl: Path, kesh, yoz, atama: str = "slayd") -> dict:
    """Har sahifa rasmga aylantiriladi -> vision OCR + rasm fragment uchun saqlanadi."""
    import fitz   # PyMuPDF
    hujjat = fitz.open(str(fayl))
    jami = len(hujjat)
    yoz(f"PDF: {fayl.name} | {jami} sahifa | vision OCR, parallel {PARALLEL}")

    rasmlar = []   # fitz thread-safe emas — rasmlarni oldin tayyorlaymiz
    for i in range(jami):
        piks = hujjat[i].get_pixmap(matrix=fitz.Matrix(RASM_MASSHTAB, RASM_MASSHTAB))
        rasmlar.append(piks.tobytes("jpeg", jpg_quality=JPEG_SIFAT))
    hujjat.close()
    yoz(f"  {jami} sahifa rasmga aylantirildi")

    natijalar: dict[int, str] = {}
    with ThreadPoolExecutor(max_workers=PARALLEL) as ex:
        # STT dagi kabi: xarajat konteksti har oqimga alohida nusxa bilan.
        # 3595 sahifalik to'plamning OCR sarfi shu tuzatishgacha daftarga
        # tushmagan edi.
        fut = {ex.submit(contextvars.copy_context().run,
                         _sahifa_ocr, r, i, jami, kesh, yoz): i
               for i, r in enumerate(rasmlar)}
        for f in as_completed(fut):
            natijalar[fut[f]] = f.result()

    bolaklar = []
    for i in sorted(natijalar):
        matn = natijalar[i].strip()
        if not matn:
            continue
        bolaklar.append({"joy": f"{i + 1}-{atama}", "matn": matn, "sahifa": i})
    return {"bolaklar": bolaklar, "davomiylik": None,
            "rasmlar": [(i, r, "image/jpeg") for i, r in enumerate(rasmlar)]}


# ---------------------------------------------------------------- jadval / hujjat / web

def _katak(v) -> str:
    """Excel katagi → markdown jadval katagi.

    `|` ekranlanadi (aks holda ustunlar siljib ketadi), sana va son
    o'qiladigan ko'rinishga keltiriladi, haddan uzun matn qisqartiriladi.
    """
    if v is None:
        return ""
    if isinstance(v, bool):
        return "ha" if v else "yo'q"
    if isinstance(v, datetime):
        if (v.hour, v.minute, v.second) == (0, 0, 0):
            return v.date().isoformat()
        return v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")):
            return ""                       # NaN / cheksizlik — bo'sh katak
        v = round(v, 6)
        v = int(v) if v == int(v) else v
    s = str(v).replace("\n", " ").replace("|", "\\|").strip()
    return s[:KATAK_UZUNLIK] + "…" if len(s) > KATAK_UZUNLIK else s


def _ustun_tozala(qatorlar: list[list[str]]) -> list[list[str]]:
    """Butunlay bo'sh ustun va qatorlarni tashlab yuboradi.

    Excel varag'i formatlash izlari tufayli o'nlab bo'sh ustun qaytarishi
    mumkin. O'shanda markdown jadval «| | | | | |» devoriga aylanadi:
    manba oynasida ham o'qib bo'lmaydi, modelga ham foyda bermaydi —
    embedding esa mazmun o'rniga tinish belgilariga o'tirib qoladi.
    """
    keng = max((len(q) for q in qatorlar), default=0)
    kerak = [i for i in range(keng)
             if any(i < len(q) and q[i] for q in qatorlar)]
    if not kerak:
        return []
    tor = []
    for q in qatorlar:
        yangi = [q[i] if i < len(q) else "" for i in kerak]
        if any(yangi):
            tor.append(yangi)
    return tor


def _varaq_qatorlari(sh, sh_formula=None) -> list[list[str]]:
    """Varaq qatorlari; bo'sh qolgan katak formula matni bilan to'ldiriladi.

    `data_only=True` faqat Excel keshlagan qiymatni beradi. Fayl skript
    bilan yasalgan yoki Excel'da hech qachon hisoblanmagan bo'lsa — kesh
    yo'q va katak bo'sh chiqadi. O'shanda `=SUM(...)` matni yoziladi:
    ko'rsatkich qanday hisoblanishi ham bilimning bir qismi.
    """
    qatorlar = [[_katak(v) for v in row] for row in sh.iter_rows(values_only=True)]
    if sh_formula is not None:
        for i, row in enumerate(sh_formula.iter_rows(values_only=True)):
            while i >= len(qatorlar):
                qatorlar.append([])
            q = qatorlar[i]
            for j, v in enumerate(row):
                if not isinstance(v, str) or not v.startswith("="):
                    continue
                while len(q) <= j:
                    q.append("")
                if not q[j]:
                    q[j] = _katak(v)
    return _ustun_tozala(qatorlar)


def jadval_oqi(fayl: Path, yoz) -> dict:
    """xlsx → har varaq uchun bitta markdown jadval."""
    import openpyxl
    wb = openpyxl.load_workbook(str(fayl), data_only=True)
    wb_formula = None
    if fayl.stat().st_size <= FORMULA_MAKS_MB * 1024 * 1024:
        try:
            wb_formula = openpyxl.load_workbook(str(fayl), data_only=False)
        except Exception as e:                      # noqa: BLE001 — ixtiyoriy qadam
            yoz(f"  formulalar o'qilmadi ({e}) — faqat qiymatlar")
    bolaklar = []
    for sh in wb.worksheets:
        sh_formula = (wb_formula[sh.title]
                      if wb_formula is not None and sh.title in wb_formula.sheetnames
                      else None)
        qatorlar = _varaq_qatorlari(sh, sh_formula)
        if not qatorlar:
            continue
        keng = len(qatorlar[0])
        md = [f"## Varaq: {sh.title}", "",
              "| " + " | ".join(qatorlar[0]) + " |",
              "|" + "---|" * keng]
        md += ["| " + " | ".join(q) + " |" for q in qatorlar[1:]]
        bolaklar.append({"joy": f"varaq: {sh.title}", "matn": "\n".join(md)})
        yoz(f"  varaq '{sh.title}': {len(qatorlar)} qator × {keng} ustun")
    return {"bolaklar": bolaklar, "davomiylik": None, "rasmlar": []}


def hujjat_oqi(fayl: Path, yoz) -> dict:
    s = fayl.suffix.lower()
    if s == ".docx":
        from docx import Document
        doc = Document(str(fayl))
        qismlar = [p.text for p in doc.paragraphs if p.text.strip()]
        for jadval in doc.tables:
            for qator in jadval.rows:
                qismlar.append(" | ".join(k.text.strip() for k in qator.cells))
        matn = "\n\n".join(qismlar)
    else:
        matn = fayl.read_text(encoding="utf-8", errors="replace")
    yoz(f"Hujjat: {fayl.name} | {len(matn)} belgi")
    return {"bolaklar": matn_bolaklar(matn), "davomiylik": None, "rasmlar": []}


BRAUZER_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36")


def web_oqi(url: str, yoz) -> dict:
    """Veb-sahifadan toza matn (trafilatura), bo'lmasa oddiy HTML tozalash."""
    import httpx
    yoz(f"Veb-sahifa olinmoqda: {url[:120]}")
    # ba'zi saytlar (Wikipedia va h.k.) "bot" ko'rinishidagi UA ni 403 qiladi
    bosh = {"User-Agent": BRAUZER_UA,
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "uz,ru;q=0.9,en;q=0.8"}
    r = httpx.get(url, timeout=60, follow_redirects=True, headers=bosh)
    if r.status_code in (403, 429):
        yoz(f"  {r.status_code} — boshqa User-Agent bilan qayta urinilmoqda")
        r = httpx.get(url, timeout=60, follow_redirects=True,
                      headers={**bosh, "User-Agent":
                               "DigitalTwin/1.0 (+https://twin.uz; bilim bazasi uchun)"})
    r.raise_for_status()
    html = r.text
    matn, sarlavha = "", ""
    try:
        import trafilatura
        matn = trafilatura.extract(html, include_comments=False,
                                   include_tables=True) or ""
        meta = trafilatura.extract_metadata(html)
        sarlavha = (getattr(meta, "title", "") or "") if meta else ""
    except Exception as e:
        yoz(f"  trafilatura ishlamadi ({str(e)[:80]}) — oddiy tozalash")
    if not matn.strip():
        toza = re.sub(r"(?is)<(script|style|nav|footer|header)[^>]*>.*?</\1>", " ", html)
        toza = re.sub(r"(?s)<[^>]+>", "\n", toza)
        matn = re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", toza)).strip()
        m = re.search(r"(?is)<title[^>]*>(.*?)</title>", html)
        sarlavha = (m.group(1).strip() if m else "")
    if len(matn.strip()) < 200:
        raise RuntimeError("sahifadan matn ajratilmadi (juda qisqa)")
    yoz(f"  {len(matn)} belgi matn ajratildi")
    return {"bolaklar": matn_bolaklar(matn), "davomiylik": None, "rasmlar": [],
            "sarlavha": sarlavha.strip()[:150]}
