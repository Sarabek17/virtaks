# -*- coding: utf-8 -*-
"""Diagnostika savol bankini Excel'dan bazaga yuklaydi (CJM + EJM).

    python -m platforma.savol_yukla --sinov     # faqat o'qiydi, bazaga tegmaydi
    python -m platforma.savol_yukla             # `diag_savollar` ga yozadi

Manba fayllar: `materiallar/diagnostika/CJM yangi.xlsx`, `EJM yangi.xlsx`.

IDEMPOTENT: `UNIQUE (tur, versiya, bosqich, blok, tartib)` bo'yicha upsert —
skript necha marta ishlatilsa ham savol soni o'zgarmaydi. Savol MATNI
o'zgargan bo'lsa yangilanadi, lekin `versiya` o'sha-o'sha qoladi; matn
mazmunan boshqa savolga aylansa `--versiya N` bilan yangi versiya yuklanadi
(eski diagnostikalar eski versiyaga bog'liq qoladi).

Varaqlarning umumiy qoidasi (ikkala faylda ham bir xil):
  * sarlavha qatorida `T/r` ustuni bor;
  * savol matni — `T/r` dan keyingi ustunda;
  * blok nomi — `T/r` dan oldingi ustunda (faqat blok boshida yoziladi);
  * bosqich sarlavhasi — "1-TANILISH" ko'rinishidagi alohida qator yoki
    "TANILISH – INSTAGRAM" ko'rinishidagi blok sarlavhasi.
"""
import argparse
import re
import sys
from pathlib import Path

from . import pg
from .sozlama import log

MANBA = Path(__file__).resolve().parent.parent / "materiallar" / "diagnostika"
CJM_FAYL = MANBA / "CJM yangi.xlsx"
EJM_FAYL = MANBA / "EJM yangi.xlsx"

# CJM'ning bu bosqichi faqat ofisli bizneslarga tegishli (tashqi reklama,
# fasad, avtoturargoh, kirish qismi) — onlayn biznesda foiz maxrajidan
# chiqariladi. Egasining qarori: kiritiladi, lekin o'chirsa bo'ladi.
IXTIYORIY_BOSQICH = {"Muhit"}

# Bosqich sarlavhasi: "1-TANILISH", "5-USHLAB QOLISH VA CHIQISH"
_BOSQICH = re.compile(r"^\s*(\d+)\s*[-–—]\s*([^\d].{2,})$")
# Blok sarlavhasi: "TANILISH – INSTAGRAM", "1-SOTUV – CALL OPERATOR"
_BLOK_SARLAVHA = re.compile(r"^(.*?)[–—](.+)$")
_RAQAM = re.compile(r"^\d+([.,]\d+)?$")


def _s(q) -> str:
    return "" if q is None else str(q).strip()


# Qisqartmalar kichik harfga tushmasin: "HR boshqaruvi", "Xodimlar AI ..."
QISQARTMA = ("HR", "AI", "KPI", "CRM", "NPS", "SMM", "PR", "SSP", "EJM", "CJM")


def _bosqich_tozala(matn: str) -> str:
    """'2-SARALASH VA ISHONCH' -> 'Saralash va ishonch'."""
    m = _BOSQICH.match(matn)
    xom = (m.group(2) if m else matn).strip(" -–—:")
    xom = re.sub(r"\s+", " ", xom)
    if not xom:
        return ""
    natija = xom[:1].upper() + xom[1:].lower()
    # Qisqartmalarni SO'Z sifatida tiklaymiz ("hr" -> "HR"); so'z ICHIDAGI
    # tasodifiy moslikka tegmaymiz ("zaif" dagi "ai" o'sha holicha qoladi).
    katta = {x.lower(): x for x in QISQARTMA}
    return " ".join(katta.get(s.lower(), s) for s in natija.split(" "))


def _varaq_oqi(ws, tur: str, bosqich_asos: str = "") -> list[dict]:
    """Bitta varaqni umumiy qoida bo'yicha o'qiydi."""
    qatorlar = list(ws.iter_rows(values_only=True))
    c_tr = None
    bosh = 0
    for i, r in enumerate(qatorlar[:8]):
        for j, c in enumerate(r):
            if _s(c).lower() in ("t/r", "t\\r", "№"):
                c_tr, bosh = j, i + 1
                break
        if c_tr is not None:
            break
    if c_tr is None:
        return []
    c_savol, c_blok = c_tr + 1, max(c_tr - 1, 0)

    natija: list[dict] = []
    bosqich = bosqich_asos
    blok = ""
    tartib = 0
    for r in qatorlar[bosh:]:
        v = [_s(c) for c in r]
        while len(v) <= c_savol:
            v.append("")
        savol = v[c_savol]
        tr = v[c_tr]

        # 1) Bosqich sarlavhasi — savol yo'q, biror ustunda "N-NOM"
        if not savol:
            sarlavha = next((x for x in v if _BOSQICH.match(x)), "")
            if sarlavha:
                yangi = _bosqich_tozala(sarlavha)
                # "1-SOTUV – CALL OPERATOR" — bosqich ham, blok ham shu yerda
                m = _BLOK_SARLAVHA.match(sarlavha)
                if m:
                    yangi = _bosqich_tozala(m.group(1))
                    blok = re.sub(r"\s+", " ", m.group(2)).strip()
                    tartib = 0
                if not bosqich_asos and yangi:
                    bosqich = yangi
                continue
            # "Muhit – Tashqi reklama" (bosqichsiz blok sarlavhasi)
            qoshma = next((x for x in v if ("–" in x or "—" in x)), "")
            if qoshma:
                m = _BLOK_SARLAVHA.match(qoshma)
                if m:
                    blok = re.sub(r"\s+", " ", m.group(2)).strip()
                    tartib = 0
                continue

        # 2) Blok nomi (faqat blok boshidagi qatorda yoziladi)
        if v[c_blok] and not _RAQAM.match(v[c_blok]):
            blok = re.sub(r"\s+", " ", v[c_blok]).strip()
            tartib = 0

        # 3) Savol qatori: T/r raqam bo'lishi shart
        if savol and _RAQAM.match(tr):
            tartib += 1
            natija.append({
                "tur": tur,
                "bosqich": bosqich or "Umumiy",
                "blok": blok or "Umumiy",
                "tartib": tartib,
                "matn": re.sub(r"\s+", " ", savol).strip()[:500],
            })
    return natija


def oqi_cjm(yol: Path = CJM_FAYL) -> list[dict]:
    """CJM: har bosqich alohida varaqda (yig'ma 'CJM' varag'i o'tkazib yuboriladi)."""
    import openpyxl
    wb = openpyxl.load_workbook(yol, data_only=True)
    natija = []
    for nom in wb.sheetnames:
        if nom in ("Dashboard", "CJM"):
            continue
        natija += _varaq_oqi(wb[nom], "cjm", bosqich_asos=nom.strip())
    return natija


def oqi_ejm(yol: Path = EJM_FAYL) -> list[dict]:
    """EJM: bitta 'EJM' varag'i, ichida bosqich sarlavhalari."""
    import openpyxl
    wb = openpyxl.load_workbook(yol, data_only=True)
    return _varaq_oqi(wb["EJM"], "ejm")


def _dublni_tozala(savollar: list[dict]) -> list[dict]:
    """Bir blok ichida bir xil matn ikki marta uchrasa, ikkinchisi tashlanadi."""
    korilgan, toza = set(), []
    for s in savollar:
        kalit = (s["tur"], s["bosqich"], s["blok"], s["matn"].lower())
        if kalit in korilgan:
            continue
        korilgan.add(kalit)
        toza.append(s)
    # tartibni blok ichida qayta raqamlaymiz (tashlangan savoldan keyin teshik qolmasin)
    hisob: dict = {}
    for s in toza:
        k = (s["tur"], s["bosqich"], s["blok"])
        hisob[k] = hisob.get(k, 0) + 1
        s["tartib"] = hisob[k]
    return toza


def yukla(versiya: int = 1, sinov: bool = False) -> dict:
    savollar = _dublni_tozala(oqi_cjm() + oqi_ejm())
    if not savollar:
        raise SystemExit("savol topilmadi — manba fayllarni tekshiring")

    hisobot: dict = {}
    for s in savollar:
        hisobot.setdefault(s["tur"], {}).setdefault(s["bosqich"], 0)
        hisobot[s["tur"]][s["bosqich"]] += 1

    if sinov:
        return {"jami": len(savollar), "hisobot": hisobot, "yozildi": 0}

    yozildi = 0
    with pg.ulanish() as u, u.cursor() as k:
        for s in savollar:
            k.execute(
                """INSERT INTO diag_savollar(tur, bosqich, blok, tartib, matn,
                                             versiya, ixtiyoriy)
                   VALUES(%s,%s,%s,%s,%s,%s,%s)
                   ON CONFLICT (tur, versiya, bosqich, blok, tartib)
                   DO UPDATE SET matn = EXCLUDED.matn,
                                 ixtiyoriy = EXCLUDED.ixtiyoriy""",
                (s["tur"], s["bosqich"], s["blok"], s["tartib"], s["matn"],
                 versiya, s["bosqich"] in IXTIYORIY_BOSQICH))
            yozildi += 1
    log(f"savol banki yuklandi: {yozildi} ta (versiya {versiya})")
    return {"jami": len(savollar), "hisobot": hisobot, "yozildi": yozildi}


SQL_FAYL = Path(__file__).resolve().parent / "migratsiyalar" / "012_savollar.sql"


def _sql_qochir(q: str) -> str:
    """Postgres matn literali (bitta tirnoq ikkilanadi)."""
    return "'" + str(q).replace("'", "''") + "'"


def sql_yoz(versiya: int = 1, yol: Path = SQL_FAYL) -> dict:
    """Savol bankini MIGRATSIYA fayliga yozadi.

    Nega shunday: `materiallar/*.xlsx` deploy paketiga kirmaydi (va serverda
    Excel o'qish shart emas). Bank migratsiya bo'lib ketsa — har muhitda
    `python -m platforma.pg` bilan o'z-o'zidan yuklanadi va qayta ishga
    tushirilsa ham dubl bermaydi (ON CONFLICT).

    Savol MATNI o'zgarsa bu fayl qayta yozilmaydi — yangi `versiya` bilan
    yangi migratsiya qo'shiladi, aks holda o'tkazilgan diagnostikalarning
    foizi keyinchalik "o'z-o'zidan" o'zgarib ketardi.
    """
    savollar = _dublni_tozala(oqi_cjm() + oqi_ejm())
    if not savollar:
        raise SystemExit("savol topilmadi - manba fayllarni tekshiring")
    qatorlar = []
    for x in savollar:
        qatorlar.append(
            "  (%s,%s,%s,%d,%s,%d,%s)" % (
                _sql_qochir(x["tur"]), _sql_qochir(x["bosqich"]),
                _sql_qochir(x["blok"]), x["tartib"], _sql_qochir(x["matn"]),
                versiya,
                "true" if x["bosqich"] in IXTIYORIY_BOSQICH else "false"))
    bosh = [
        "-- Diagnostika savol banki (versiya %d) - %d savol." % (versiya, len(savollar)),
        "-- AVTOMATIK YARATILGAN: `python -m platforma.savol_yukla --sql`.",
        "-- Manba: materiallar/diagnostika/CJM yangi.xlsx, EJM yangi.xlsx.",
        "-- Qo'lda tahrirlanmaydi. Savol matni o'zgarsa - YANGI versiya bilan",
        "-- yangi migratsiya fayli qo'shiladi (eski diagnostika foizi",
        "-- o'zgarmasin).",
        "INSERT INTO diag_savollar(tur, bosqich, blok, tartib, matn,",
        "                          versiya, ixtiyoriy) VALUES",
    ]
    matn = (chr(10).join(bosh) + chr(10) +
            ("," + chr(10)).join(qatorlar) + chr(10) +
            "ON CONFLICT (tur, versiya, bosqich, blok, tartib) DO UPDATE SET" +
            chr(10) +
            "  matn = EXCLUDED.matn, ixtiyoriy = EXCLUDED.ixtiyoriy;" + chr(10))
    yol.write_text(matn, encoding="utf-8", newline=chr(10))
    return {"jami": len(savollar), "yol": str(yol)}


def main():
    p = argparse.ArgumentParser(description="Diagnostika savol bankini yuklash")
    p.add_argument("--sinov", action="store_true", help="bazaga yozmasdan tekshirish")
    p.add_argument("--versiya", type=int, default=1)
    p.add_argument("--sql", action="store_true",
                   help="bazaga emas, migratsiya fayliga yozish")
    a = p.parse_args()

    if a.sql:
        n = sql_yoz(versiya=a.versiya)
        print(f"{n['jami']} savol -> {n['yol']}")
        return
    n = yukla(versiya=a.versiya, sinov=a.sinov)
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    for tur, bosqichlar in n["hisobot"].items():
        jami = sum(bosqichlar.values())
        print(f"\n=== {tur.upper()} — {jami} savol ===")
        for b, soni in bosqichlar.items():
            print(f"   {soni:>4}  {b}")
    print(f"\nJAMI: {n['jami']} savol"
          + (" (sinov — bazaga yozilmadi)" if a.sinov else
             f", bazaga yozildi: {n['yozildi']}"))


if __name__ == "__main__":
    main()
