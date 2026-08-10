# -*- coding: utf-8 -*-
"""DOCX / TXT / MD -> oddiy matn."""
from pathlib import Path


def oqi(fayl: Path) -> dict:
    suffix = fayl.suffix.lower()
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(fayl))
        qismlar = [p.text for p in doc.paragraphs if p.text.strip()]
        for jadval in doc.tables:
            for qator in jadval.rows:
                qismlar.append(" | ".join(k.text.strip() for k in qator.cells))
        return {"matn": "\n\n".join(qismlar)}
    return {"matn": fayl.read_text(encoding="utf-8", errors="replace")}
