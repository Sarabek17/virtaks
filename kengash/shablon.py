# -*- coding: utf-8 -*-
"""Shablon xizmati — savol hujjat tuzishni so'rasa (reja, jadval, matritsa...),
kengash xulosasi asosida ma'lumotlari to'ldirilgan tayyor Excel fayl yasaydi.

Shablonlar: kengash/shablonlar/*.xlsx (asl fayllar o'zgarmaydi — nusxa to'ldiriladi).
Xato hech qachon majlisni buzmaydi — chaqiruvchi try/except bilan o'raydi.
"""
import json
import re
import time

from .agentlar import _generatsiya
from .sozlama import AGENT_MODELLAR, CHIQISH, ILDIZ, RAIS_MODELLAR, log

SHABLONLAR = ILDIZ / "shablonlar"
FAYLLAR = CHIQISH / "fayllar"          # to'ldirilgan fayllar shu yerda
KATAK_CHEGARA = 250                    # har varaqdan nechta katak LLM ga ko'rsatiladi


def royxat() -> list[str]:
    if not SHABLONLAR.is_dir():
        return []
    return sorted(f.name for f in SHABLONLAR.glob("*.xlsx"))


def _json_ol(matn: str):
    """Birinchi to'liq JSON obyektni oladi — model JSONdan keyin ortiqcha
    matn qo'shsa ham yiqilmaydi (raw_decode qoldiqni e'tiborsiz qoldiradi)."""
    return json.JSONDecoder().raw_decode(matn.strip())[0]


def mos(savol: str) -> str | None:
    """Savol shablon asosida HUJJAT TUZISHNI so'rayaptimi? -> fayl nomi yoki None."""
    r = royxat()
    if not r:
        return None
    prompt = f"""Foydalanuvchi AI kengashga savol berdi: {savol}

Qo'ldagi shablon fayllar:
{chr(10).join(f"- {n}" for n in r)}

Savol yuqoridagi shablonlardan biri asosida HUJJAT/JADVAL TUZIB BERISHNI so'rayaptimi?
(masalan: "Gantt chart tuz", "vaqt jadvali kerak", "Eyzenxauer matritsasi qilib ber",
"gipoteza yozib ber", "tashkiliy tuzilma tuz", "SSP/КПI jadvali kerak")

Qoidalar:
- Faqat savol chindan hujjat TUZIShNI so'rasa va mos shablon bo'lsa tanla.
- Oddiy ma'lumot/maslahat savoli bo'lsa ("Gantt chart nima?", "qanday sotamiz?") -> null.
- "namuna" fayllar to'ldirilgan misollar, "shablon" fayllar bo'sh — ikkalasi ham tanlanishi
  mumkin, lekin bo'sh "shablon" varianti bor bo'lsa o'shani afzal ko'r.

Javob — faqat JSON: {{"shablon": "aniq fayl nomi"}} yoki {{"shablon": null}}"""
    try:
        j = _json_ol(_generatsiya(AGENT_MODELLAR, prompt, json_rejim=True,
                                  harorat=0.0, tez=True))
        s = j.get("shablon")
        if s and s in r:
            return s
    except Exception as e:
        log(f"Shablon aniqlash xatosi: {str(e)[:80]}")
    return None


def _kataklar(wb) -> dict:
    """Har varaqdan katak inventarizatsiyasi {varaq: {manzil: qiymat}}."""
    nat = {}
    for sh in wb.worksheets:
        kk = {}
        for row in sh.iter_rows():
            for c in row:
                if c.value is not None and str(c.value).strip():
                    kk[c.coordinate] = str(c.value)[:120]
                    if len(kk) >= KATAK_CHEGARA:
                        break
            if len(kk) >= KATAK_CHEGARA:
                break
        if kk:
            nat[sh.title] = kk
    return nat


def toldir(nom: str, savol: str, xulosa: str):
    """Shablon nusxasini kengash xulosasi asosida to'ldirib, yangi fayl yo'lini qaytaradi."""
    import openpyxl
    wb = openpyxl.load_workbook(str(SHABLONLAR / nom))
    inv = _kataklar(wb)

    prompt = f"""Excel shablonini to'ldirish vazifasi.

SHABLON FAYL: {nom}
Shablondagi mavjud kataklar (varaq -> katak manzili: hozirgi qiymati):
{json.dumps(inv, ensure_ascii=False)[:9000]}

FOYDALANUVCHI SAVOLI: {savol}

KENGASH XULOSASI (to'ldirish uchun ma'lumot manbasi):
{xulosa[:7000]}

Vazifa: shablon TUZILISHINI saqlagan holda foydalanuvchi vaziyatiga mos ANIQ qiymatlar
bilan to'ldir (o'zbek tilida):
- Sarlavha/izoh kataklariga TEGMA — ular shablonning skeleti.
- Namuna-qiymatlar o'rniga foydalanuvchi vaziyatiga moslarini yoz.
- Bo'sh qolishi kerak joyni yozma. Jadval davomiga yangi qatorlar qo'shsang, mavjud
  ustun tartibini saqla (masalan A12, B12, C12...).
- Kamida 8-10 mazmunli katak to'ldir.

Javob — FAQAT JSON: {{"Varaq nomi": {{"A5": "qiymat", "B7": "qiymat"}}}}"""
    j = _json_ol(_generatsiya(RAIS_MODELLAR, prompt, json_rejim=True, harorat=0.2))

    soni = 0
    for varaq, kk in (j or {}).items():
        if varaq not in wb.sheetnames or not isinstance(kk, dict):
            continue
        sh = wb[varaq]
        for adr, q in kk.items():
            if not re.fullmatch(r"[A-Z]{1,3}\d{1,5}", str(adr).strip().upper()):
                continue
            try:
                sh[str(adr).strip().upper()] = q
                soni += 1
            except Exception:
                pass    # birlashgan (merged) katakning ichki yacheykasi — o'tkazamiz
    if not soni:
        raise RuntimeError("shablon to'ldirilmadi — LLM javobi bo'sh/yaroqsiz")

    FAYLLAR.mkdir(parents=True, exist_ok=True)
    yangi = FAYLLAR / f"{time.strftime('%Y%m%d_%H%M%S')}_{nom}"
    wb.save(str(yangi))
    log(f"Shablon to'ldirildi: {yangi.name} ({soni} katak)")
    return yangi
